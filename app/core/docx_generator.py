import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

_PROJECT_ROOT = Path(__file__).resolve().parents[2]
_OUTPUT_DIR = _PROJECT_ROOT / "output" / "generated"


def summarize_experience(experience_text: str, max_items: int = 3) -> List[str]:
    """Return a concise, role-friendly set of experience bullets."""
    if not experience_text:
        return []

    lines = [line.strip() for line in experience_text.splitlines() if line.strip()]
    if not lines:
        return []

    cleaned = []
    for line in lines:
        line = re.sub(r"\s+", " ", line)
        line = re.sub(r"^[-•*]\s*", "", line)
        if len(line) > 220:
            line = line[:217].rstrip() + "..."
        cleaned.append(line)

    return cleaned[:max_items]


def select_relevant_projects(projects: List[Dict[str, Any]], role_title: str, jd_text: str) -> List[Dict[str, Any]]:
    """Pick the best projects for the target role and trim them to a small set."""
    if not projects:
        return []

    def _normalize(token: str) -> str:
        token = re.sub(r"[^a-z0-9]+", "", token.lower())
        if token.endswith("ing") and len(token) > 4:
            token = token[:-3]
        elif token.endswith("s") and len(token) > 3:
            token = token[:-1]
        elif token.endswith("ed") and len(token) > 4:
            token = token[:-2]
        return token

    jd_terms = {_normalize(token) for token in re.split(r"[^a-z0-9]+", f"{role_title} {jd_text}") if token}
    jd_terms = {term for term in jd_terms if len(term) > 2}

    scored = []
    for project in projects:
        title = str(project.get("title", "")).lower()
        description = str(project.get("description", "")).lower()
        technologies = " ".join(project.get("technologies", [])).lower()
        content = f"{title} {description} {technologies}"
        content_terms = {_normalize(token) for token in re.split(r"[^a-z0-9]+", content) if token}

        overlap = sum(1 for term in jd_terms if term and term in content_terms)
        title_overlap = sum(3 for term in jd_terms if term and term in {_normalize(token) for token in re.split(r"[^a-z0-9]+", title) if token})
        tech_bonus = sum(2 for term in ["python", "xgboost", "prophet", "ml", "api", "forecast", "deploy", "backend", "data"] if term in content)
        score = overlap * 2 + title_overlap + tech_bonus
        scored.append((score, project))

    scored.sort(key=lambda item: item[0], reverse=True)
    selected = [project for _, project in scored[:2]]
    return selected


def _ensure_output_dir(path: str | os.PathLike[str]) -> str:
    path_obj = Path(path)
    if not path_obj.is_absolute():
        path_obj = _PROJECT_ROOT / path_obj
    path_obj.parent.mkdir(parents=True, exist_ok=True)
    return str(path_obj)


def _add_heading(doc: Document, text: str) -> None:
    p = doc.add_heading(text, level=1)
    p.runs[0].font.size = Pt(12)


def generate_resume_docx(
    profile: Dict[str, Any],
    tailored_bullets: List[str],
    company_name: str,
    role_title: str,
    output_dir: str = _OUTPUT_DIR,
) -> str:
    doc = Document()

    # Compact page margins
    for section in doc.sections:
        section.top_margin = Pt(36)
        section.bottom_margin = Pt(36)
        section.left_margin = Pt(54)
        section.right_margin = Pt(54)

    # Name
    name_para = doc.add_paragraph()
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name_para.add_run(profile.get("name", "Your Name"))
    run.bold = True
    run.font.size = Pt(18)

    # Contact line
    contacts = [
        profile.get("email", ""),
        profile.get("phone", ""),
        profile.get("linkedin", ""),
        profile.get("github", ""),
    ]
    contact_line = "  |  ".join(c for c in contacts if c)
    cp = doc.add_paragraph(contact_line)
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()

    # Summary / tailored opening
    if tailored_bullets:
        _add_heading(doc, "Summary")
        summary_text = tailored_bullets[0]
        if len(summary_text) > 220:
            summary_text = summary_text[:217].rstrip() + "..."
        doc.add_paragraph(summary_text)

    # Tailored highlights
    if len(tailored_bullets) > 1:
        _add_heading(doc, f"Highlights")
        for bullet in tailored_bullets[1:4]:
            doc.add_paragraph(bullet, style="List Bullet")

    # Experience (trimmed to keep the resume concise)
    experience = profile.get("experience", "").strip()
    if experience:
        _add_heading(doc, "Experience")
        for line in summarize_experience(experience, max_items=3):
            doc.add_paragraph(line, style="List Bullet")

    # Projects (selected and condensed to stay one-page friendly)
    projects_text = profile.get("projects", "").strip()
    project_list: List[Dict[str, Any]] = profile.get("project_list", [])
    selected_projects = select_relevant_projects(project_list, role_title, profile.get("jd_text", "")) if project_list else []
    if selected_projects or projects_text:
        _add_heading(doc, "Selected Projects")
        for project in selected_projects[:3]:
            title = project.get("title", "")
            desc = project.get("description", "")
            techs = ", ".join(project.get("technologies", []))
            if title:
                p = doc.add_paragraph()
                p.add_run(title).bold = True
            if desc:
                desc_text = desc if len(desc) <= 220 else desc[:217].rstrip() + "..."
                doc.add_paragraph(desc_text)
            if techs:
                doc.add_paragraph(f"Technologies: {techs}")

    # Education (kept short)
    education = profile.get("education", "").strip()
    if education:
        _add_heading(doc, "Education")
        for line in education.splitlines()[:3]:
            if line.strip():
                doc.add_paragraph(line.strip())

    # Skills (kept compact)
    skills: List[str] = profile.get("skills", [])
    if skills:
        _add_heading(doc, "Skills")
        doc.add_paragraph(", ".join(skills[:12]))

    safe_company = company_name.replace(" ", "_")
    date_str = datetime.now().strftime("%Y%m%d")
    output_dir_path = Path(output_dir) if isinstance(output_dir, (str, os.PathLike)) else _OUTPUT_DIR
    if not output_dir_path.is_absolute():
        output_dir_path = _PROJECT_ROOT / output_dir_path
    output_dir_path.mkdir(parents=True, exist_ok=True)
    path = _ensure_output_dir(output_dir_path / f"resume_{safe_company}_{date_str}.docx")
    doc.save(path)
    return path


def generate_cover_letter_docx(
    profile: Dict[str, Any],
    company_name: str,
    role_title: str,
    cover_letter_text: str,
    output_dir: str = _OUTPUT_DIR,
) -> str:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Pt(54)
        section.bottom_margin = Pt(54)
        section.left_margin = Pt(72)
        section.right_margin = Pt(72)

    # Sender header
    name_p = doc.add_paragraph()
    name_p.add_run(profile.get("name", "Your Name")).bold = True
    contacts = [profile.get("email", ""), profile.get("phone", ""), profile.get("linkedin", "")]
    doc.add_paragraph("  |  ".join(c for c in contacts if c))
    doc.add_paragraph(datetime.now().strftime("%B %d, %Y"))
    doc.add_paragraph()

    # Addressee
    doc.add_paragraph(f"Hiring Manager\n{company_name}")
    doc.add_paragraph()

    # Subject line
    subj = doc.add_paragraph()
    subj.add_run(f"Re: Application for {role_title}").bold = True
    doc.add_paragraph()

    # Body paragraphs (split on double newlines)
    for para in cover_letter_text.strip().split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())

    # Closing
    doc.add_paragraph()
    doc.add_paragraph("Sincerely,")
    doc.add_paragraph()
    doc.add_paragraph(profile.get("name", ""))

    safe_company = company_name.replace(" ", "_")
    date_str = datetime.now().strftime("%Y%m%d")
    output_dir_path = Path(output_dir) if isinstance(output_dir, (str, os.PathLike)) else _OUTPUT_DIR
    if not output_dir_path.is_absolute():
        output_dir_path = _PROJECT_ROOT / output_dir_path
    output_dir_path.mkdir(parents=True, exist_ok=True)
    path = _ensure_output_dir(output_dir_path / f"cover_letter_{safe_company}_{date_str}.docx")
    doc.save(path)
    return path
